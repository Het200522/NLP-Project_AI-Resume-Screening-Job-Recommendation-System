"""
Extracts raw text from uploaded resume/JD files (PDF, DOCX, TXT).
Falls back to OCR (Tesseract) for scanned/image-only PDFs.
"""
import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    """Raised when a file's text cannot be extracted."""


def extract_text(file_path: str | Path) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".txt":
        return _extract_txt(path)

    raise ExtractionError(f"Unsupported file type: {suffix}")


def _extract_pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ExtractionError("PyMuPDF (fitz) is not installed") from e

    try:
        text_parts = []
        with fitz.open(str(path)) as doc:
            if doc.is_encrypted:
                raise ExtractionError("PDF is password-protected and cannot be read")
            for page in doc:
                text_parts.append(page.get_text())
        text = "\n".join(text_parts).strip()
        if text:
            return text
        logger.info("No embedded text in PDF, attempting OCR on %s", path)
        return _ocr_pdf(path)
    except ExtractionError:
        raise
    except Exception as e:
        logger.exception("PDF extraction failed for %s", path)
        raise ExtractionError(f"Failed to read PDF: {e}") from e


def _ocr_pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
        import pytesseract
        from PIL import Image
    except ImportError as e:
        raise ExtractionError(
            "OCR requires pytesseract and Pillow. "
            "Install them with: pip install pytesseract Pillow"
        ) from e

    try:
        text_parts = []
        with fitz.open(str(path)) as doc:
            for page_num, page in enumerate(doc, start=1):
                pix = page.get_pixmap(dpi=300)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                page_text = pytesseract.image_to_string(img)
                if page_text.strip():
                    text_parts.append(page_text)
        text = "\n".join(text_parts).strip()
        if not text:
            raise ExtractionError(
                "OCR could not extract any text from the scanned PDF"
            )
        logger.info("OCR extracted %d characters from %s", len(text), path)
        return text
    except ExtractionError:
        raise
    except Exception as e:
        logger.exception("OCR extraction failed for %s", path)
        raise ExtractionError(f"OCR failed: {e}") from e


def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as e:
        raise ExtractionError("python-docx is not installed") from e

    try:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        text = "\n".join(parts).strip()
        if not text:
            raise ExtractionError("No extractable text found in DOCX")
        return text
    except ExtractionError:
        raise
    except Exception as e:
        logger.exception("DOCX extraction failed for %s", path)
        raise ExtractionError(f"Failed to read DOCX: {e}") from e


def _extract_txt(path: Path) -> str:
    try:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        if not text:
            raise ExtractionError("Text file is empty")
        return text
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to read text file: {e}") from e
